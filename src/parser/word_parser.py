# -*- coding: utf-8 -*-
"""
Word文档解析器
基于 python-docx 实现 .docx 文本和表格提取
"""
from .base_parser import BaseParser, ParseResult
from ..utils.logger import get_logger

logger = get_logger(__name__)


class WordParser(BaseParser):
    """Word (.docx) 文档解析器"""

    supported_extensions = ["docx"]
    parser_name = "docx"

    def parse(self, file_path: str, **kwargs) -> ParseResult:
        """
        解析Word文档，提取文本和表格内容

        Args:
            file_path: .docx文件路径
            **kwargs:
                extract_tables: 是否提取表格（默认True）
                extract_headers: 是否提取页眉页脚（默认False）
                extract_metadata: 是否提取文档属性（默认True）

        Returns:
            ParseResult
        """
        try:
            from docx import Document
        except ImportError:
            return ParseResult(
                success=False,
                file_path=file_path,
                file_type=self.parser_name,
                error="python-docx未安装，请执行: pip install python-docx",
            )

        extract_tables = kwargs.get("extract_tables", True)
        extract_headers = kwargs.get("extract_headers", False)
        extract_meta = kwargs.get("extract_metadata", True)

        doc = Document(file_path)

        # 提取段落
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:  # 过滤空行
                paragraphs.append(text)

        full_text = "\n".join(paragraphs)

        # 提取表格
        tables = []
        if extract_tables:
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                if table_data:
                    tables.append(table_data)

        # 提取页眉页脚
        if extract_headers:
            for section in doc.sections:
                header_text = ""
                footer_text = ""
                if section.header:
                    header_text = "\n".join(
                        p.text for p in section.header.paragraphs if p.text.strip()
                    )
                if section.footer:
                    footer_text = "\n".join(
                        p.text for p in section.footer.paragraphs if p.text.strip()
                    )
                if header_text:
                    paragraphs.insert(0, f"[页眉] {header_text}")
                    full_text = f"[页眉] {header_text}\n" + full_text
                if footer_text:
                    paragraphs.append(f"[页脚] {footer_text}")
                    full_text += f"\n[页脚] {footer_text}"

        # 提取文档属性
        metadata = {}
        if extract_meta:
            props = doc.core_properties
            for attr in ["author", "category", "comments", "created", "identifier",
                         "keywords", "language", "modified", "subject", "title", "version"]:
                value = getattr(props, attr, None)
                if value:
                    metadata[attr] = str(value)

        metadata["paragraphs_count"] = len(paragraphs)
        metadata["tables_count"] = len(tables)

        logger.info(f"Word解析完成: {file_path} | {len(paragraphs)}段落, {len(tables)}表格")

        return ParseResult(
            success=True,
            file_path=file_path,
            file_type=self.parser_name,
            text=full_text,
            paragraphs=paragraphs,
            tables=tables,
            metadata=metadata,
        )
